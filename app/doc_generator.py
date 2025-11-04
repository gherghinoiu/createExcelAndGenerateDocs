# app/doc_generator.py - REFACTORED to use docxtpl
import pandas as pd
import re
from pathlib import Path
from docxtpl import DocxTemplate, RichText
from docx.shared import Pt
from docx.enum.text import WD_COLOR_INDEX
from app.utils.config import (
    TEMPLATE_1_FILE,
    TEMPLATE_2_FILE,
    GENERATED_DOCS_DIR
)

# --- Configuration ---
PLACEHOLDER_MAP = {
    'apel': 'Apel',
    'nume_proiect': 'Nume proiect',
    'nume_autoritate': 'Denumire autoritate contractantă',
    'ofertant': 'Ofertant',
    'nr_anunt_sicap': 'Nr. anunt SICAP',
    'numar_contract': 'Număr contract',
    'data_semnarii': 'Data semnării contractului',
    'valoare_estimata_fara_tva': 'Valoare estimata',
    'valoare_contract_fara_tva': 'Valoare cumparare directa',
    'beneficiari_reali': 'Beneficiari reali',
    'beneficiari_reali_url_pnrr': 'Beneficiari reali URL',
    'seap_url': 'seap_url',
    'detalii_achizitie_url_pnrr': 'Detalii achizitie URL PNRR',
}

EMPTY_VALUE_REPLACEMENT = "[A SE COMPLETA DE OFITER]"

# --- Helper Functions ---

def get_jalon_tinta(row):
    """
    Finds the 'I-number' in the 'Apel' string and maps it
    to the corresponding Jalon / Țintă number.
    """
    JALON_MAP = {
        'I5': '281',
        'I8': '284',
        'I9': '285',
        'I10': '287',
    }
    
    apel_text = row.get('Apel')
    if pd.isna(apel_text):
        return None
    
    match = re.search(r'\b(I5|I8|I9|I10)\b', str(apel_text))
    if match:
        i_code = match.group(1)
        return JALON_MAP[i_code]
    else:
        return None


def _clean_value_to_float(raw_text):
    """
    Converts a raw string value like '140.000,50' or None
    into a float (e.g., 140000.50).
    Returns 0.0 if the value is empty or invalid.
    """
    if pd.isna(raw_text):
        return 0.0
    if isinstance(raw_text, (int, float)):
        return float(raw_text)
    try:
        text = str(raw_text)
        text = text.replace('.', '')
        text = text.replace(',', '.')
        text = re.sub(r"[^0-9.]", "", text)
        if not text:
            return 0.0
        return float(text)
    except (ValueError, TypeError):
        return 0.0


def get_articol_de_lege(row):
    """
    Determines the correct legal article text based on the SICAP ID
    and the contract value.
    """
    sicap_id = str(row.get('Nr. anunt SICAP', '')).strip().upper()
    alin_text = ""
    litera_text = ""
    
    if sicap_id.startswith(('DA', 'DAN', 'ADV')):
        alin_text = "Alin (7) În cazul achiziţiei directe, autoritatea contractantă:"
        valoare = _clean_value_to_float(row.get('Valoare cumparare directa'))
        
        if valoare <= 9000:
            litera_text = " d) are dreptul de a plăti direct, pe baza angajamentului legal, fără acceptarea prealabilă a unei oferte, dacă valoarea estimată a achiziţiei este mai mică de 9.000 lei, fără TVA."
        elif 9000 < valoare <= 140000:
            litera_text = " c) are dreptul de a achiziţiona pe baza unei singure oferte dacă valoarea estimată a achiziţiei este mai mică sau egală cu 140.000 lei, fără TVA, pentru produse şi servicii, respectiv 300.000 lei, fără TVA, pentru lucrări;"
        elif 140000 < valoare <= 200000:
            litera_text = " b) are obligaţia de a consulta minimum trei operatori economici pentru achiziţiile a căror valoare estimată este mai mare de 140.000 lei, fără TVA, pentru produse şi servicii, respectiv 300.000 lei, fără TVA, pentru lucrări, dar mai mică sau egală cu valoarea menţionată la lit. a); dacă în urma consultării autoritatea contractantă primeşte doar o ofertă valabilă din punctul de vedere al cerinţelor solicitate, achiziţia poate fi realizată;"
        elif valoare > 200000:
            litera_text = " a) are obligaţia de a utiliza catalogul electronic pus la dispoziţie de SEAP sau de a publica un anunţ într-o secţiune dedicată a website-ului propriu sau al SEAP, însoţit de descrierea produselor, serviciilor sau a lucrărilor care urmează a fi achiziţionate, pentru achiziţiile a căror valoare estimată este mai mare de 200.000 lei, fără TVA, pentru produse şi servicii, respectiv 560.000 lei, fără TVA, pentru lucrări;"
    
    elif sicap_id.startswith(('CN', 'CAN')):
        alin_text = " alin. (1)"
    
    elif sicap_id.startswith(('SCN', 'SCNA')):
        alin_text = " alin. (2)"
    
    else:
        return None
    
    return (alin_text + "\n" + litera_text)


def sanitize_filename(text):
    """Clean text for safe filenames."""
    if pd.isna(text) or text == "":
        return "NA"
    return re.sub(r'[\\/*?:"<>|]', '_', str(text).strip())


def get_unique_filename(row, template_type, output_dir):
    """Generate a unique filename for the output document."""
    sicap_id = sanitize_filename(row.get('Nr. anunt SICAP'))
    numar_contract = sanitize_filename(row.get('Număr contract'))
    nume_autoritate = sanitize_filename(row.get('Denumire autoritate contractantă'))
    
    base_name = f"{sicap_id}_{template_type}_{numar_contract}_{nume_autoritate}"
    final_path = output_dir / f"{base_name}.docx"
    
    counter = 1
    while final_path.exists():
        final_path = output_dir / f"{base_name}_{counter}.docx"
        counter += 1
    
    return final_path

def create_url_richtext(doc, url_str):
    """
    Helper function to create a RichText hyperlink with consistent formatting.
    
    :param doc: The DocxTemplate document object
    :param url_str: The URL string to convert to a hyperlink
    :return: RichText object with formatted hyperlink
    """
    rt = RichText()
    rt.add(url_str, url_id=doc.build_url_id(url_str), 
           font='Trebuchet MS', size=22)
    return rt

def add_custom_placeholders(row, context):
    """
    Add custom computed placeholders to the context dictionary.
    This includes logic-derived values like articol_de_lege and nr_jalon_tinta.
    """
    doc = context.get('_docx')
    
    # Add the legal article text
    article_text = get_articol_de_lege(row)
    context['articol_de_lege'] = article_text if article_text else EMPTY_VALUE_REPLACEMENT
    
    # Add the jalon/tinta number
    jalon_text = get_jalon_tinta(row)
    context['nr_jalon_tinta'] = jalon_text if jalon_text else EMPTY_VALUE_REPLACEMENT
    
    # Handle URLs with RichText for clickable links
    url_pnrr = row.get('Beneficiari reali URL')
    context['beneficiari_reali_url_pnrr'] = (
        create_url_richtext(doc, str(url_pnrr).strip()) 
        if pd.notna(url_pnrr) and str(url_pnrr).strip() 
        else EMPTY_VALUE_REPLACEMENT
    )
    
    seap_url = row.get('seap_url')
    context['seap_url'] = (
        create_url_richtext(doc, str(seap_url).strip())
        if pd.notna(seap_url) and str(seap_url).strip()
        else EMPTY_VALUE_REPLACEMENT
    )
    
    detalii_url = row.get('Detalii achizitie URL PNRR')
    context['detalii_achizitie_url_pnrr'] = (
        create_url_richtext(doc, str(detalii_url).strip())
        if pd.notna(detalii_url) and str(detalii_url).strip()
        else EMPTY_VALUE_REPLACEMENT
    )
    
    return context


def build_context_from_row(row, template_doc):
    """
    Build the context dictionary from a DataFrame row.
    Maps Excel columns to template placeholders.
    """
    context = {'_docx': template_doc}  # Store doc reference for URL building
    
    # Map standard placeholders
    for placeholder, col_name in PLACEHOLDER_MAP.items():
        value = row.get(col_name)
        if pd.isna(value) or value == "":
            context[placeholder] = EMPTY_VALUE_REPLACEMENT
        else:
            context[placeholder] = str(value)
    
    # Add custom computed placeholders
    context = add_custom_placeholders(row, context)
    
    # Remove the _docx helper from context before rendering
    context.pop('_docx', None)
    
    return context


def generate_document_from_template(template_path, row, output_path):
    """
    Generate a single document from a template and row data.
    Uses docxtpl for robust placeholder replacement.
    """
    try:
        # Load the template
        doc = DocxTemplate(template_path)
        
        # Build context dictionary
        context = build_context_from_row(row, doc)
        
        # Render the template with the context
        doc.render(context)
        
        # Save the document
        doc.save(output_path)
        return True, None
    
    except Exception as e:
        return False, str(e)


# --- Main Orchestrator Function ---

def run_document_generation(excel_files: list):
    """
    Main function to run the entire document generation process.
    'excel_files' is a list of Path objects to process.
    """
    print("--- Step 5: Generating Documents (with docxtpl) ---")
    
    all_dfs = []
    for file_path in excel_files:
        if not file_path.exists():
            print(f" > Warning: Input file not found: {file_path}")
            continue
        try:
            all_dfs.append(pd.read_excel(file_path))
            print(f" > Loaded: {file_path.name}")
        except Exception as e:
            print(f" > Error loading {file_path.name}: {e}")
    
    if not all_dfs:
        print(" > Error: No valid Excel files were loaded. Stopping.")
        return False
    
    df = pd.concat(all_dfs, ignore_index=True)
    print(f" > Loaded a total of {len(df)} rows to process.")
    
    # Ensure output directory exists
    GENERATED_DOCS_DIR.mkdir(exist_ok=True)
    
    success_count = 0
    failed_count = 0
    
    for index, row in df.iterrows():
        sicap_id = row.get('Nr. anunt SICAP', f'Row {index+1}')
        print(f"\n Processing {index + 1}/{len(df)}: {sicap_id}")
        
        # --- Process Template 1 (LV) ---
        save_path_1 = get_unique_filename(row, 'LV', GENERATED_DOCS_DIR)
        success, error = generate_document_from_template(TEMPLATE_1_FILE, row, save_path_1)
        
        if success:
            print(f" > Saved LV: {save_path_1.name}")
            success_count += 1
        else:
            print(f" > FAILED to generate LV doc: {error}")
            failed_count += 1
        
        # --- Process Template 2 (RV) ---
        save_path_2 = get_unique_filename(row, 'RV', GENERATED_DOCS_DIR)
        success, error = generate_document_from_template(TEMPLATE_2_FILE, row, save_path_2)
        
        if success:
            print(f" > Saved RV: {save_path_2.name}")
            success_count += 1
        else:
            print(f" > FAILED to generate RV doc: {error}")
            failed_count += 1
    
    print(f"\n--- Document Generation Complete ---")
    print(f" > Successfully generated: {success_count} documents")
    print(f" > Failed: {failed_count} documents")
    
    return True
